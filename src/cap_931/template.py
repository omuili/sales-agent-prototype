import re
from datetime import date
from pathlib import Path

import docx
from docx.document import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .layout import PAGE_SIDE_MARGIN_IN

TEMPLATE_PATH = Path(__file__).with_name("assets") / "perscholas_template.docx"
BODY_FONT = "Helvetica Neue"
LINK_COLOR = "1F5FAD"
MUTED = RGBColor(0x66, 0x66, 0x66)
BRAND = RGBColor(0x61, 0x19, 0x40)  


_PLAIN = str.maketrans({"\u2010": "-", "\u2011": "-", "\u00a0": " ", "\u202f": " ", "\u2009": " "})


def clean(text: str) -> str:
    return text.translate(_PLAIN)


def _set_fonts(rpr_owner, name: str) -> None:
    """Set the font for every script slot so no character falls back to another typeface."""
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)

KICKER, TITLE, LOGO, BYLINE = 0, 1, 2, 3


def _set_text(paragraph, text: str) -> None:
    """Replace a paragraph's text but keep the formatting of its first run."""
    runs = [r for r in paragraph.runs if r.text]
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def open_template(kicker: str, title: str, learner: str, when: date | None = None, keep_toc: bool = False) -> Document:
    """Open the template with its cover filled in and the placeholder sections removed.

    keep_toc keeps the cover's Table of Contents (for long reports); Word refreshes it when the file is opened.
    """
    doc = docx.Document(str(TEMPLATE_PATH))
    paragraphs = doc.paragraphs
    _set_text(paragraphs[KICKER], kicker)
    _set_text(paragraphs[TITLE], title)
    byline = [r for r in paragraphs[BYLINE].runs if r.text]
    byline[0].text = f"{learner or 'Learner'}\n"
    byline[1].text = f"{when or date.today():%B %d, %Y}"

    body = doc.element.body
    cover_end = paragraphs[BYLINE]._p
    started = False
    for child in list(body):
        if child is cover_end:
            started = True
            continue
        if not started or child.tag == qn("w:sectPr"):
            continue
        if keep_toc and child.tag == qn("w:sdt"):
            # List only Heading 2 (sections) and Heading 3 (steps) so the contents fit on the cover.
            for instr in child.iter(qn("w:instrText")):
                if "TOC" in (instr.text or ""):
                    instr.text = ' TOC \\h \\u \\z \\t "Heading 2,1,Heading 3,2" '
            continue
        body.remove(child)
    if keep_toc:
        _update_fields_on_open(doc)
        page_break(doc)
    return doc


def open_page() -> Document:
    """The template with an empty body: keeps its header, logo, footer and styles, but no cover page."""
    doc = docx.Document(str(TEMPLATE_PATH))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = doc.sections[0]
    # The cover uses a different first-page header; a single page should use the regular one.
    section.different_first_page_header_footer = False
    section.left_margin = section.right_margin = Inches(PAGE_SIDE_MARGIN_IN)
    # The template numbers pages from 0 (its cover is page 0); a single page should be page 1.
    for numbering in section._sectPr.findall(qn("w:pgNumType")):
        section._sectPr.remove(numbering)
    return doc


def _update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        flag = OxmlElement("w:updateFields")
        flag.set(qn("w:val"), "true")
        settings.append(flag)


def page_break(doc: Document) -> None:
    run = doc.add_paragraph(style="normal").add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def heading(doc: Document, text: str, level: int = 2):
    return doc.add_paragraph(clean(text), style=f"Heading {min(max(level, 1), 6)}")


def label(doc: Document, text: str, size: float = 11):
    """A compact maroon section label (smaller than the template's headings, for dense pages)."""
    para = doc.add_paragraph(style="normal")
    para.paragraph_format.space_before = Pt(7)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.keep_with_next = True
    run(para, text, size=size, bold=True, color=BRAND)
    return para


def paragraph(doc: Document, text: str = "", size: float = 10.5, italic: bool = False, color: RGBColor | None = None):
    para = doc.add_paragraph(style="normal")
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.0 if size < 10 else 1.15
    if text:
        add_inline(para, text, size=size, italic=italic, color=color)
    return para


def bullet(doc: Document, text: str = "", size: float = 10.5, level: int = 0):
    """A bullet paragraph (the template has no list style, so use a hanging indent)."""
    para = doc.add_paragraph(style="normal")
    fmt = para.paragraph_format
    fmt.left_indent = Pt(14 + 14 * level)
    fmt.first_line_indent = Pt(-10)
    # The template's body text uses 1.5 line spacing; dense pages (small text) need it tighter.
    fmt.line_spacing = 1.0 if size < 10 else 1.15
    fmt.space_after = Pt(1.5 if size < 10 else 3)
    run(para, "•  " if level == 0 else "◦  ", size=size)
    if text:
        add_inline(para, text, size=size)
    return para


def run(para, text: str, size: float = 10.5, bold: bool = False, italic: bool = False,
        color: RGBColor | None = None, mono: bool = False):
    r = para.add_run(clean(text))
    _set_fonts(r._r, "Courier New" if mono else BODY_FONT)
    r.font.size = Pt(size - 1 if mono else size)
    r.bold, r.italic = bold or None, italic or None
    if color:
        r.font.color.rgb = color
    return r


def hyperlink(para, url: str, text: str, size: float = 10.5) -> None:
    rel_id = para.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    r = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), BODY_FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    for el in (fonts, color, underline, sz):
        props.append(el)
    r.append(props)
    t = OxmlElement("w:t")
    t.text = clean(text)
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    para._p.append(link)


def highlight(para, text: str, size: float = 10.5):
    r = run(para, text, size=size, bold=True)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


LINK = re.compile(r"\[(\[[^\]]+\]|[^\]]+)\]\(([^)\s]+)\)")
INLINE_SPLIT = re.compile(
    r"(\*\*[^*]+\*\*"                            # **bold**
    r"|`[^`]+`"                                  # `code`
    r"|\[(?:\[[^\]]+\]|[^\]]+)\]\([^)\s]+\)"     # [text](url) or [[n]](url)
    r"|\*[^*\s][^*]*\*|_[^_\s][^_]*_)"           # *italic* or _italic_
)


def add_inline(para, text: str, size: float = 10.5, italic: bool = False, color: RGBColor | None = None) -> None:
    """Add text with Markdown-style **bold**, *italic*, `code` and [links](url)."""
    for piece in INLINE_SPLIT.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run(para, piece[2:-2], size=size, bold=True, italic=italic, color=color)
        elif piece.startswith("`") and piece.endswith("`"):
            run(para, piece[1:-1], size=size, mono=True)
        elif piece.startswith("[") and (m := LINK.fullmatch(piece)):
            if m.group(2).startswith("http"):
                hyperlink(para, m.group(2), m.group(1), size=size)
            else:
                run(para, m.group(1), size=size, italic=italic, color=color)
        elif len(piece) > 2 and piece[0] in "*_" and piece[-1] == piece[0]:
            run(para, piece[1:-1], size=size, italic=True, color=color)
        else:
            run(para, piece, size=size, italic=italic, color=color)


def table(doc: Document, rows: list[list[str]], size: float = 9.5, widths_in: list[float] | None = None) -> None:
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    _borders(tbl)
    if widths_in:
        tbl.autofit = False
        for row in tbl.rows:
            for cell, width in zip(row.cells, widths_in):
                cell.width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            para = cell.paragraphs[0]
            para.style = doc.styles["normal"]
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, value.replace("**", "") if i == 0 else value, size=size)
            if i == 0:
                for r in para.runs:
                    r.bold = True
                _shade(cell, "E8EEF7")
    doc.add_paragraph(style="normal")


def _borders(tbl) -> None:
    props = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    props.append(borders)


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def code_block(doc: Document, text: str) -> None:
    for line in text.rstrip("\n").split("\n"):
        para = doc.add_paragraph(style="normal")
        fmt = para.paragraph_format
        fmt.space_after = fmt.space_before = Pt(0)
        fmt.left_indent = Pt(10)
        fmt.line_spacing = 1.0
        run(para, line or " ", size=8.5, mono=True)
    doc.add_paragraph(style="normal")


LIST_ITEM = re.compile(r"^\s*([-*]|\d+\.)\s+")


def add_markdown(doc: Document, md: str, heading_offset: int = 0, size: float = 10.5) -> None:
    """Render a subset of Markdown (headings, bullets, numbered lists, tables, code, quotes, paragraphs)."""
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
        elif stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, "\n".join(block))
            i += 1
        elif m := re.match(r"^(#{1,6})\s+(.*)", stripped):
            heading(doc, m.group(2), len(m.group(1)) + heading_offset)
            i += 1
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows)
            table(doc, [r + [""] * (width - len(r)) for r in rows])
        elif m := re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line):
            text = m.group(3)
            i += 1
            # Indented lines that aren't new list items continue the current item.
            while i < len(lines) and re.match(r"^\s+\S", lines[i]) and not LIST_ITEM.match(lines[i]):
                text += " " + lines[i].strip()
                i += 1
            para = bullet(doc, size=size, level=len(m.group(1)) // 2)
            if m.group(2)[0].isdigit():
                para.runs[0].text = f"{m.group(2)}  "
            add_inline(para, text, size=size)
        elif stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            para = paragraph(doc, " ".join(quote), size=size, italic=True, color=MUTED)
            para.paragraph_format.left_indent = Pt(14)
        else:
            text = [stripped]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^\s*([-*]\s|\d+\.\s|#|\||>|```)", lines[i]):
                text.append(lines[i].strip())
                i += 1
            paragraph(doc, " ".join(text), size=size)
