import io
from pathlib import Path

import docx
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED = (".pdf", ".pptx", ".docx", ".txt", ".md")
MAX_CHARS = 15000


def extract_deck_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".pptx":
        slides = []
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            parts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
            slides.append(f"[Slide {i}] " + " ".join(p for p in parts if p.strip()))
        text = "\n".join(slides)
    elif suffix == ".docx":
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    elif suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type {suffix!r}; use one of {', '.join(SUPPORTED)}")
    return text.strip()[:MAX_CHARS]
