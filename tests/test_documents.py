from docx.oxml.ns import qn

from cap_931 import template as t


def test_template_cover_is_filled_and_placeholders_removed():
    doc = t.open_template("Account Brief", "Acme Bank", "Jane Learner")
    text = [p.text for p in doc.paragraphs]
    assert text[:2] == ["Account Brief", "Acme Bank"]
    assert text[3].startswith("Jane Learner\n")
    assert not any("[" in line for line in text)  # no [Title]/[Learner Name]/[Date] or lab placeholders
    assert doc.element.body.find(qn("w:sdt")) is None  # TOC dropped for the one-pager


def test_open_page_has_no_cover():
    doc = t.open_page()
    assert doc.paragraphs == [] and doc.element.body.find(qn("w:sdt")) is None


def test_markdown_lists_join_wrapped_lines():
    doc = t.open_template("k", "T", "L")
    t.add_markdown(doc, "1. First line\n   continues\n- bullet")
    assert [p.text for p in doc.paragraphs[4:]] == ["1.  First line continues", "•  bullet"]
