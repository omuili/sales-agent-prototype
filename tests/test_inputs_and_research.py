import io

import docx
from pptx import Presentation

from cap_931.research import edgar
from cap_931.research.deck import extract_deck_text
from cap_931.research.fetch import internal_links, site_name
from cap_931.research.gather import focus_text, short_name
from cap_931.schemas import CompanyProfile
from cap_931.validation import validate_inputs


def test_valid_inputs_are_normalized():
    inputs, errors = validate_inputs("FlowLake", "capitalone.com", "Unifies data at lower cost.",
                                     competitors="snowflake.com, https://www.databricks.com\nsnowflake.com")
    assert errors == []
    assert inputs.company_url == "https://capitalone.com"
    assert inputs.competitor_urls == ["https://snowflake.com", "https://www.databricks.com"]


def test_invalid_inputs_report_every_problem():
    inputs, errors = validate_inputs("", "not a url", "short", competitors="capitalone.com, bad url")
    assert inputs is None
    joined = " ".join(errors)
    assert "Product name is required" in joined
    assert "not a valid website" in joined
    assert "Value proposition is required" in joined
    assert "'bad url'" in joined


def test_competitor_cannot_be_the_prospect():
    _, errors = validate_inputs("X", "https://acme.com", "A valid value proposition.", competitors="www.acme.com")
    assert errors == ["Competitor 'www.acme.com' is the same site as the prospect."]


def test_too_many_competitors():
    _, errors = validate_inputs("X", "acme.com", "A valid value proposition.",
                                competitors="\n".join(f"c{i}.com" for i in range(6)))
    assert any("at most 5" in e for e in errors)


def test_site_name_prefers_structured_data_then_title():
    ld = '<html><head><script type="application/ld+json">{"@type": "Organization", "name": "Acme Bank"}</script>' \
         '<title>Home | Something</title></head></html>'
    assert site_name(ld, "https://acmebank.com") == "Acme Bank"
    assert site_name("<html><head><title>Credit Cards | Capital One</title></head></html>",
                     "https://capitalone.com") == "Capital One"


def test_short_name_uses_domain_brand():
    assert short_name(CompanyProfile(name="Snowflake AI Data Cloud", domain="snowflake.com", url="")) == "Snowflake"
    assert short_name(CompanyProfile(name="Capital One", domain="capitalone.com", url="")) == "Capital One"


def test_internal_links_grouped_by_page_type():
    page = ('<a href="/about-us">About</a><a href="/leadership">Leaders</a><a href="https://other.com/news">x</a>'
            '<a href="https://investor.acme.com/">IR</a><a href="/careers/">Jobs</a>')
    links = internal_links(page, "https://acme.com")
    assert links["leadership"] == ["https://acme.com/leadership"]
    assert links["about"] == ["https://acme.com/about-us"]
    assert links["investors"] == ["https://investor.acme.com"]
    assert links["careers"] == ["https://acme.com/careers"]
    assert "news" not in links  # other domains are ignored


def test_focus_text_keeps_paragraphs_about_the_company():
    text = "\n".join(["filler " * 50] * 5 + ["Acme Bank signed a deal."])
    out = focus_text(text, ["acme bank"], 400)
    assert out.startswith("Acme Bank signed a deal.")
    assert len(out) <= 400


def test_edgar_section_skips_table_of_contents():
    toc = "Item 1. Business 5 Item 1A. Risk Factors 20 "
    body = "Item 1. Business " + "We are a bank. " * 300 + "Item 1A. Risk Factors ..."
    assert edgar._section(toc + body, r"item\s*1\.?\s*business", r"item\s*1a", 100).startswith("We are a bank.")


def test_edgar_name_normalization():
    assert edgar._normalize("Capital One Financial Corp.") == "capital one financial"
    assert edgar._normalize("AT&T Inc.") == "at and t"


def test_deck_extraction_pptx_docx_txt():
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text = "FlowLake"
    slide.placeholders[1].text = "Governed lakehouse"
    buf = io.BytesIO()
    pres.save(buf)
    assert "Governed lakehouse" in extract_deck_text("deck.pptx", buf.getvalue())

    document = docx.Document()
    document.add_paragraph("Cuts costs 40%")
    buf = io.BytesIO()
    document.save(buf)
    assert extract_deck_text("one.docx", buf.getvalue()) == "Cuts costs 40%"
    assert extract_deck_text("notes.txt", b"hello") == "hello"


def test_iso_date_handles_rfc_and_iso():
    from cap_931.research.search import iso_date

    assert iso_date("Thu, 06 Nov 2025 17:00:00 GMT") == "2025-11-06"
    assert iso_date("2026-10-16T02:59:54+00:00") == "2026-10-16"
    assert iso_date("") == "" and iso_date("garbage") == ""
