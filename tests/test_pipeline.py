import io

import docx
import pytest

from cap_931 import agents
from cap_931.pipeline import OutOfScopeError, run_pipeline
from cap_931.render import cited_ids, to_docx, to_markdown
from cap_931.schemas import CompanyProfile, Intake

from .conftest import EVIDENCE, FakeLLM


def test_full_pipeline_runs_every_agent_and_prunes_unsupported(sales_input, fake_gather):
    llm = FakeLLM()
    pager = run_pipeline(sales_input, llm)

    assert llm.calls[0] == "Intake"
    assert {"StrategySection", "CompetitorSection", "LeadershipSection", "FinancialSection", "OnePageBrief",
            "VerifierOutput"} <= set(llm.calls)
    texts = [c.text for c in pager.strategy.key_points]
    assert texts == ["Acme is moving analytics to a lakehouse"]  # unsupported + bad-citation claims removed
    assert pager.strategy.key_points[0].verdict == "supported"
    assert any("removed 2" in w for w in pager.warnings)
    assert pager.leadership.leaders[0].name == "Jane Roe"


def test_out_of_scope_request_stops_before_research(sales_input, fake_gather):
    llm = FakeLLM(in_scope=False)
    with pytest.raises(OutOfScopeError, match="Not a sales request"):
        run_pipeline(sales_input, llm)
    assert llm.calls == ["Intake"]


def test_pipeline_without_verifier_keeps_claims(sales_input, fake_gather):
    llm = FakeLLM()
    pager = run_pipeline(sales_input, llm, verify=False)
    assert "VerifierOutput" not in llm.calls
    assert len(pager.strategy.key_points) == 3


def test_markdown_and_one_page_docx(sales_input, fake_gather):
    pager = run_pipeline(sales_input, FakeLLM())
    md = to_markdown(pager)
    for heading in ("### Executive Summary", "### Key Contacts", "### Talking Points", "# Detailed Research",
                    "## Competitor Mentions", "## Leadership", "## Product & Financial Strategy", "## Sources"):
        assert heading in md
    assert "[[1]](https://acmebank.com/tech)" in md
    assert '"Data is our product"' in md  # quotes stay in the detailed research
    assert "### Company Strategy" not in md  # empty brief sections are skipped, not padded
    assert cited_ids(pager) == [1, 2, 3]

    document = docx.Document(io.BytesIO(to_docx(pager)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert text.startswith("Account Brief: Acme Bank")
    assert "Olanrewaju Muili" in text and "Jane Roe" in text
    assert "[" not in text.split("\n")[0] and "Overview" not in text  # no template placeholders
    assert not document.sections[0].different_first_page_header_footer  # single page, regular header
    assert sum(1 for p in document.paragraphs if 'w:type="page"' in p._p.xml) == 0
    links = {r.target_ref for r in document.part.rels.values() if r.is_external}
    assert {"https://acmebank.com/tech", "https://sec.gov/acme-10k"} <= links


def test_writer_tightens_an_over_budget_brief(sales_input, fake_gather):
    llm = FakeLLM(brief_words=600)
    run_pipeline(sales_input, llm, verify=False)
    assert llm.brief_calls == 3  # first draft + two tightening passes (the fake never gets shorter)


def _ctx(competitors):
    return agents.Context(
        inputs=None, intake=Intake(in_scope=True, category="x", search_keywords=[], product_summary=""),
        company=CompanyProfile(name="Acme Bank", domain="acmebank.com", url="https://acmebank.com"),
        competitors=competitors, evidence=EVIDENCE,
    )


def test_competitor_hits_are_found_deterministically():
    ctx = _ctx([CompanyProfile(name="Snowflake", domain="snowflake.com", url="https://snowflake.com"),
                CompanyProfile(name="Databricks", domain="databricks.com", url="https://databricks.com")])
    assert agents.competitor_hits(ctx) == {"Snowflake": [1], "Databricks": []}


def test_competitor_agent_skips_llm_when_nothing_mentions_competitors():
    ctx = _ctx([CompanyProfile(name="Teradata", domain="teradata.com", url="https://teradata.com")])
    llm = FakeLLM()
    section = agents.run_competitors(llm, ctx)
    assert llm.calls == []
    assert [(m.competitor, m.mentioned) for m in section.mentions] == [("Teradata", False)]


def test_evidence_budget_skips_oversized_blocks():
    text = agents.format_evidence(EVIDENCE, budget=400)
    assert 'id="1"' in text and 'id="3"' in text
    assert agents.format_evidence([]) == "<evidence>No evidence was found.</evidence>"
